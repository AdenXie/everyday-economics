"""Extract matching original ABS CPI and WPI quarterly index series.

Requires openpyxl. Raw ABS workbooks stay in data/sources; this writes only
the small, browser-ready data/quarterly.json snapshot.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = ROOT / "data" / "source-manifest.json"
OUTPUT_PATH = ROOT / "data" / "quarterly.json"


def load_series(source: dict[str, Any]) -> tuple[dict[str, float], dict[str, float | None]]:
    workbook_path = ROOT / source["workbook"]
    if not workbook_path.is_file():
        raise FileNotFoundError(f"ABS source workbook not found: {workbook_path}")
    expected_hash = source.get("sha256")
    if expected_hash:
        actual_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
        if actual_hash.lower() != expected_hash.lower():
            raise ValueError(f"Source hash mismatch for {workbook_path.name}: {actual_hash}")

    book = load_workbook(workbook_path, data_only=True, read_only=True)
    if "Index" not in book.sheetnames or "Data1" not in book.sheetnames:
        raise ValueError(f"Unexpected workbook layout: {workbook_path.name}")

    index_sheet = book["Index"]
    metadata = {}
    for row in index_sheet.iter_rows(min_row=12, values_only=True):
        if row[4] == source["indexSeriesId"]:
            metadata = {
                "description": str(row[0] or ""),
                "seriesType": row[3],
                "seriesStart": row[5],
                "seriesEnd": row[6],
                "observations": row[7],
                "unit": row[8],
                "dataType": row[9],
                "frequency": row[10],
            }
            break

    if not metadata:
        raise ValueError(f"Series ID {source['indexSeriesId']} is missing from Index sheet")
    if metadata["seriesType"] != source["seriesType"]:
        raise ValueError(f"Unexpected adjustment type for {source['indexSeriesId']}: {metadata['seriesType']}")
    if " ".join(metadata["description"].split()) != " ".join(source["seriesDescription"].split()):
        raise ValueError(f"Series description changed for {source['indexSeriesId']}: {metadata['description']}")
    if metadata["unit"] != "Index Numbers" or metadata["dataType"] != "INDEX" or metadata["frequency"] != "Quarter":
        raise ValueError(f"Expected an original quarterly index series, got {metadata}")

    data_sheet = book["Data1"]
    rows = data_sheet.iter_rows(values_only=True)
    headers = next(rows)
    units = next(rows)
    series_types = next(rows)
    data_types = next(rows)
    frequencies = next(rows)
    collection_months = next(rows)
    starts = next(rows)
    ends = next(rows)
    counts = next(rows)
    series_ids = next(rows)

    try:
        index_col = list(series_ids).index(source["indexSeriesId"])
        change_col = list(series_ids).index(source["quarterChangeSeriesId"])
    except ValueError as error:
        raise ValueError(f"Series IDs are missing from Data1 in {workbook_path.name}") from error

    if " ".join(str(headers[index_col]).split()) != " ".join(source["seriesDescription"].split()):
        raise ValueError(f"Index description mismatch for {source['indexSeriesId']}")
    if units[index_col] != "Index Numbers" or series_types[index_col] != "Original":
        raise ValueError(f"Expected original index values for {source['indexSeriesId']}")
    if data_types[index_col] != "INDEX" or frequencies[index_col] != "Quarter":
        raise ValueError(f"Unexpected series metadata for {source['indexSeriesId']}")
    if "Percentage Change from Previous Period" not in str(headers[change_col]) and "Percentage Change from Previous Quarter" not in str(headers[change_col]):
        raise ValueError(f"Quarter-change description mismatch for {source['quarterChangeSeriesId']}")
    if units[change_col] != "Percent" or series_types[change_col] != "Original":
        raise ValueError(f"Expected original quarterly percent changes for {source['quarterChangeSeriesId']}")

    indexes: dict[str, float] = {}
    changes: dict[str, float | None] = {}
    for row in rows:
        raw_period = row[0]
        if not isinstance(raw_period, (datetime, date)):
            continue
        quarter_by_month = {3: 1, 6: 2, 9: 3, 12: 4}
        if raw_period.month not in quarter_by_month:
            raise ValueError(f"Non-quarter date in quarterly source: {raw_period}")
        period = f"{raw_period.year}-Q{quarter_by_month[raw_period.month]}"
        raw_index = row[index_col]
        raw_change = row[change_col]
        if raw_index is None:
            continue
        index_value = float(raw_index)
        if not math.isfinite(index_value) or index_value <= 0:
            raise ValueError(f"Invalid index value for {period}: {raw_index}")
        indexes[period] = index_value
        changes[period] = None if raw_change is None else float(raw_change)

    if not indexes:
        raise ValueError(f"No observations found for {source['indexSeriesId']}")
    if len(indexes) != metadata["observations"]:
        raise ValueError(f"Observation count mismatch for {source['indexSeriesId']}: {len(indexes)} vs {metadata['observations']}")
    first_period = next(iter(indexes))
    last_period = next(reversed(indexes))
    expected_start = metadata["seriesStart"]
    expected_end = metadata["seriesEnd"]
    if not isinstance(expected_start, datetime) or not isinstance(expected_end, datetime):
        raise ValueError(f"Missing date bounds for {source['indexSeriesId']}")
    expected_start_quarter = f"{expected_start.year}-Q{quarter_by_month[expected_start.month]}"
    expected_end_quarter = f"{expected_end.year}-Q{quarter_by_month[expected_end.month]}"
    if first_period != expected_start_quarter or last_period != expected_end_quarter:
        raise ValueError(f"Date bounds disagree for {source['indexSeriesId']}: {first_period}..{last_period}")
    return indexes, changes


def quarter_number(period: str) -> int:
    year, quarter = period.split("-Q")
    return int(year) * 4 + int(quarter)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract the official ABS CPI/WPI quarterly snapshot")
    parser.add_argument("--snapshot-date", help="Date the static snapshot was prepared (YYYY-MM-DD)")
    args = parser.parse_args()

    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    if args.snapshot_date:
        date.fromisoformat(args.snapshot_date)
        manifest["preparedOn"] = args.snapshot_date

    cpi, cpi_change = load_series(manifest["cpi"])
    wpi, wpi_change = load_series(manifest["wpi"])
    shared_periods = sorted(set(cpi) & set(wpi), key=quarter_number)
    if not shared_periods:
        raise ValueError("CPI and WPI have no common quarterly observations")

    records = []
    for period in shared_periods:
        records.append({
            "period": period,
            "cpiIndex": cpi[period],
            "wpiIndex": wpi[period],
            "cpiQuarterChangePct": cpi_change.get(period),
            "wpiQuarterChangePct": wpi_change.get(period),
        })

    common_start = max(min(cpi, key=quarter_number), min(wpi, key=quarter_number), key=quarter_number)
    common_end = min(max(cpi, key=quarter_number), max(wpi, key=quarter_number), key=quarter_number)
    expected_common_count = quarter_number(common_end) - quarter_number(common_start) + 1
    if shared_periods[0] != common_start or shared_periods[-1] != common_end or len(shared_periods) != expected_common_count:
        raise ValueError("CPI and WPI do not provide every quarter in their common date range")
    for previous, current in zip(shared_periods, shared_periods[1:]):
        if quarter_number(current) - quarter_number(previous) != 1:
            raise ValueError(f"Quarter gap between {previous} and {current}")

    snapshot = {
        "snapshotDate": manifest["preparedOn"],
        "firstQuarter": shared_periods[0],
        "latestQuarter": shared_periods[-1],
        "observationCount": len(records),
        "series": {
            "cpi": {
                "publication": manifest["cpi"]["publication"],
                "releaseDate": manifest["cpi"]["releaseDate"],
                "table": manifest["cpi"]["table"],
                "pageUrl": manifest["cpi"]["pageUrl"],
                "seriesId": manifest["cpi"]["indexSeriesId"],
                "seriesDescription": manifest["cpi"]["seriesDescription"],
                "seriesType": manifest["cpi"]["seriesType"],
                "indexReferencePeriod": manifest["cpi"]["indexReferencePeriod"],
                "sourceWorkbook": manifest["cpi"]["workbook"],
                "sourceSha256": manifest["cpi"].get("sha256"),
                "scope": "Australia; All groups CPI; weighted average of eight capital cities",
            },
            "wpi": {
                "publication": manifest["wpi"]["publication"],
                "releaseDate": manifest["wpi"]["releaseDate"],
                "table": manifest["wpi"]["table"],
                "pageUrl": manifest["wpi"]["pageUrl"],
                "seriesId": manifest["wpi"]["indexSeriesId"],
                "seriesDescription": manifest["wpi"]["seriesDescription"],
                "seriesType": manifest["wpi"]["seriesType"],
                "indexReferencePeriod": manifest["wpi"]["indexReferencePeriod"],
                "sourceWorkbook": manifest["wpi"]["workbook"],
                "sourceSha256": manifest["wpi"].get("sha256"),
                "scope": "Australia; private and public sectors; all industries; total hourly rates of pay excluding bonuses",
            },
        },
        "observations": records,
    }

    OUTPUT_PATH.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUTPUT_PATH.relative_to(ROOT)}: {len(records)} aligned quarters, {shared_periods[0]} to {shared_periods[-1]}")


if __name__ == "__main__":
    main()
